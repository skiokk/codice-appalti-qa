from collections.abc import Callable
from datetime import datetime
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from appalti_qa.evaluation.dataset import DatasetError, EvalItem, load_dataset
from appalti_qa.evaluation.xlsx import (
    SHEET_ORIGINAL,
    SHEET_REVIEW,
    export_xlsx,
    import_xlsx,
    read_review,
    review_import,
)


def _headers(ws) -> dict[str, int]:
    return {cell.value: cell.column for cell in ws[1]}


def _review(path: Path, edits: dict[str, dict[str, object]]) -> None:
    """Fill the review sheet: {item id: {header: value}}."""
    workbook = load_workbook(path)
    ws = workbook[SHEET_REVIEW]
    columns = _headers(ws)
    rows = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    for item_id, values in edits.items():
        row = rows.get(item_id)
        if row is None:
            row = ws.max_row + 1
            ws.cell(row=row, column=1, value=item_id)
        for header, value in values.items():
            ws.cell(row=row, column=columns[header], value=value)
    workbook.save(path)


@pytest.fixture
def workbook_path(tmp_path: Path, items: list[EvalItem]) -> Path:
    path = tmp_path / "review" / "revisione.xlsx"
    export_xlsx(items, path)
    return path


def test_export_layout(workbook_path: Path, items: list[EvalItem]) -> None:
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == [SHEET_REVIEW, "Istruzioni", SHEET_ORIGINAL]
    assert workbook[SHEET_ORIGINAL].sheet_state == "hidden"
    ws = workbook[SHEET_REVIEW]
    headers = _headers(ws)
    assert list(headers) == [
        "ID",
        "Tipo",
        "Categoria",
        "Vigenza",
        "Domanda",
        "Risposta attesa",
        "Riferimenti",
        "Estratto",
        "Esito",
        "Correzione",
        "Difficoltà",
        "Stato",
        "Note",
    ]
    assert ws.freeze_panes == "B2"
    assert ws["A1"].font.bold
    assert ws.column_dimensions["E"].width == 50
    assert ws.cell(row=2, column=headers["Riferimenti"]).value == "50/1/b"
    assert ws.cell(row=5, column=headers["Riferimenti"]).value == "1/1, 50/1"
    assert ws.cell(row=2, column=headers["Vigenza"]).value == "2025-01-01"
    assert ws.cell(row=2, column=headers["Domanda"]).alignment.wrap_text
    (validation,) = ws.data_validations.dataValidation
    assert validation.formula1 == '"OK,Da correggere,Scartare"'
    assert str(validation.sqref) == "I2:I5"


def test_export_optional_columns(
    tmp_path: Path, items: list[EvalItem], make_answer: Callable
) -> None:
    path = tmp_path / "with_answers.xlsx"
    answers = {"Q001": make_answer(items[0], "=risposta che inizia con uguale", ["50/1"])}
    export_xlsx(items, path, answers=answers, unit_texts={"Q001": "[Art. 50, comma 1]\n1. ..."})
    ws = load_workbook(path)[SHEET_REVIEW]
    headers = _headers(ws)
    assert ws.cell(row=2, column=headers["Testo vigente"]).value.startswith("[Art. 50")
    assert ws.cell(row=2, column=headers["Risposta del sistema"]).value.startswith("=risposta")
    assert ws.cell(row=2, column=headers["Riferimenti citati"]).value == "50/1"
    assert ws.cell(row=3, column=headers["Risposta del sistema"]).value in (None, "")


def test_import_round_trip(workbook_path: Path, dataset_path: Path) -> None:
    before = dataset_path.read_text(encoding="utf-8")
    _review(
        workbook_path,
        {
            "Q001": {"Esito": "OK"},
            "Q002": {"Esito": "Da correggere", "Correzione": "Verificare la data del correttivo."},
            "Q003": {"Esito": "scartare", "Correzione": "Domanda ambigua."},
            "Q004": {
                "Esito": "Da correggere",
                "Risposta attesa": "Risposta riscritta dal revisore.",
                "Riferimenti": "1/1; 50/1/b",
                "Vigenza": datetime(2025, 2, 1),
            },
        },
    )
    items = import_xlsx(workbook_path, dataset_path)
    by_id = {item.id: item for item in items}
    assert by_id["Q001"].status == "verificata"
    assert by_id["Q002"].status == "bozza"  # was corretta; a problem was reported
    assert (
        by_id["Q002"].notes
        == "Changed by D.Lgs. 209/2024.\nRevisione: Verificare la data del correttivo."
    )
    assert by_id["Q003"].status == "scartata"
    assert by_id["Q004"].status == "corretta"
    assert by_id["Q004"].expected_answer == "Risposta riscritta dal revisore."
    assert by_id["Q004"].expected_refs == ["1/1", "50/1/b"]
    assert str(by_id["Q004"].vigenza) == "2025-02-01"
    assert load_dataset(dataset_path) == items
    after = dataset_path.read_text(encoding="utf-8")
    assert after.startswith("# Synthetic evaluation set for the tests.\n")
    assert "# --- temporal questions ---" in after
    assert [entry["id"] for entry in yaml.safe_load(after)] == ["Q001", "Q002", "Q003", "Q004"]
    assert len(after.splitlines()) == len(before.splitlines())

    # importing the same workbook again changes nothing
    result = review_import(workbook_path, dataset_path)
    assert result.changed == []
    assert dataset_path.read_text(encoding="utf-8") == after


def test_ok_keeps_corretta(workbook_path: Path, dataset_path: Path) -> None:
    _review(workbook_path, {"Q002": {"Esito": "OK"}})
    result = review_import(workbook_path, dataset_path)
    assert result.changed == []


def test_dry_run_writes_nothing(workbook_path: Path, dataset_path: Path) -> None:
    before = dataset_path.read_text(encoding="utf-8")
    _review(workbook_path, {"Q001": {"Esito": "Scartare"}})
    result = review_import(workbook_path, dataset_path, write=False)
    assert result.changes == ["Q001: verificata -> scartata"]
    assert dataset_path.read_text(encoding="utf-8") == before


def test_rows_without_outcome_and_unknown_ids(workbook_path: Path, dataset_path: Path) -> None:
    _review(
        workbook_path,
        {"Q001": {"Domanda": "Domanda modificata ma senza esito?"}, "Q999": {"Esito": "OK"}},
    )
    result = review_import(workbook_path, dataset_path)
    assert result.changed == []
    assert result.warnings == [
        "Q001: edited but Esito is empty, row ignored",
        "row 6: unknown id Q999, ignored",
    ]


@pytest.mark.parametrize(
    ("edits", "message"),
    [
        ({"Esito": "Forse"}, "invalid Esito 'Forse'"),
        ({"Esito": "OK", "Riferimenti": "art. 50"}, "invalid key"),
        ({"Esito": "OK", "Tipo": "boh"}, "kind"),
        ({"Esito": "OK", "Vigenza": "ieri"}, "vigenza"),
    ],
)
def test_invalid_review_leaves_the_file_untouched(
    workbook_path: Path, dataset_path: Path, edits: dict, message: str
) -> None:
    before = dataset_path.read_text(encoding="utf-8")
    _review(workbook_path, {"Q001": edits})
    with pytest.raises(DatasetError) as error:
        import_xlsx(workbook_path, dataset_path)
    assert any(message in problem for problem in error.value.problems), error.value.problems
    assert dataset_path.read_text(encoding="utf-8") == before


def test_conflict_with_later_dataset_edit(workbook_path: Path, dataset_path: Path) -> None:
    text = dataset_path.read_text(encoding="utf-8")
    dataset_path.write_text(text.replace("è ammesso sotto", "è consentito sotto"), "utf-8")
    _review(workbook_path, {"Q001": {"Esito": "OK", "Risposta attesa": "Nuova risposta."}})
    with pytest.raises(DatasetError, match="changed in the dataset after the export"):
        import_xlsx(workbook_path, dataset_path)


def test_sheet_without_hidden_originals(workbook_path: Path, dataset_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    del workbook[SHEET_ORIGINAL]
    workbook.save(workbook_path)
    _review(workbook_path, {"Q004": {"Esito": "OK", "Domanda": "Domanda corretta?"}})
    assert read_review(workbook_path).originals is None
    q004 = import_xlsx(workbook_path, dataset_path)[3]
    assert (q004.question, q004.status) == ("Domanda corretta?", "corretta")


def test_changed_items_are_validated(workbook_path: Path, dataset_path: Path) -> None:
    _review(workbook_path, {"Q003": {"Esito": "OK", "Riferimenti": "50/1"}})
    with pytest.raises(DatasetError, match="unanswerable questions must have no expected_refs"):
        import_xlsx(workbook_path, dataset_path)


def test_missing_columns(tmp_path: Path, dataset_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "plain.xlsx"
    workbook = Workbook()
    workbook.active.append(["ID", "Domanda"])
    workbook.save(path)
    with pytest.raises(DatasetError, match="missing column"):
        import_xlsx(path, dataset_path)
