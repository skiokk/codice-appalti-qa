import json
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from datetime import date
from pathlib import Path

import psycopg
import pytest
from openpyxl import load_workbook

from appalti_qa.config import get_settings
from appalti_qa.evaluation import cli as eval_cli
from appalti_qa.evaluation import cli_main
from appalti_qa.evaluation.dataset import EvalItem, load_dataset
from appalti_qa.evaluation.judge import OverlapJudge
from appalti_qa.evaluation.xlsx import SHEET_REVIEW
from appalti_qa.models import Answer, Hit


@pytest.fixture
def no_db(monkeypatch: pytest.MonkeyPatch) -> None:
    """Commands that need a connection get None; retrieval goes to the stub retriever."""

    @contextmanager
    def open_db() -> Iterator[None]:
        yield None

    monkeypatch.setattr(eval_cli, "_open_db", open_db)


def _run_file(out: Path) -> Path:
    """The results file of the only answer run written to `out`."""
    (path,) = (p for p in out.glob("*.json") if p.name != "latest.json")
    return path


@pytest.fixture
def stub_backends(
    monkeypatch: pytest.MonkeyPatch, no_db: None, stub_retriever: Callable, stub_answer: Callable
) -> list[str]:
    modes: list[str] = []

    def search_retriever(mode: str, conn: object, settings: object) -> Callable:
        modes.append(mode)
        return stub_retriever

    def answer_backend(settings, provider_name, model, retrieve, *, k, mode, delay_s):
        return stub_answer, provider_name, model or "fake"

    monkeypatch.setattr(eval_cli, "_search_retriever", search_retriever)
    monkeypatch.setattr(eval_cli, "_answer_backend", answer_backend)
    return modes


def test_validate(dataset_path: Path, capsys: pytest.CaptureFixture) -> None:
    assert cli_main(["validate", "--dataset", str(dataset_path)]) == 0
    assert "4 questions, 2 verified; static checks passed" in capsys.readouterr().out


def test_validate_reports_problems(tmp_path: Path, capsys: pytest.CaptureFixture) -> None:
    path = tmp_path / "bad.yaml"
    path.write_text(
        "- id: Q001\n  question: q\n  kind: multi\n  category: c\n  vigenza: '2025-01-01'\n"
        "  expected_answer: a\n  expected_refs: ['50/1']\n  source_excerpt: x\n"
        "  difficulty: easy\n",
        encoding="utf-8",
    )
    assert cli_main(["validate", "--dataset", str(path)]) == 1
    err = capsys.readouterr().err
    assert "Q001: multi questions need at least two expected_refs" in err
    assert "1 problem(s) found." in err
    assert cli_main(["validate", "--dataset", str(tmp_path / "missing.yaml")]) == 1


@pytest.mark.db
def test_validate_with_db(
    monkeypatch: pytest.MonkeyPatch,
    seeded_conn: psycopg.Connection,
    dataset_path: Path,
    capsys: pytest.CaptureFixture,
) -> None:
    @contextmanager
    def open_db() -> Iterator[psycopg.Connection]:
        yield seeded_conn

    monkeypatch.setattr(eval_cli, "_open_db", open_db)
    argv = ["validate", "--db", "--dataset", str(dataset_path)]
    assert cli_main(argv) == 0
    assert "static and corpus checks passed" in capsys.readouterr().out
    # Q002 quotes the text in force before the 2024-12-31 amendment
    text = dataset_path.read_text(encoding="utf-8")
    dataset_path.write_text(text.replace('"2024-06-01"', '"2025-06-01"'), encoding="utf-8")
    assert cli_main(argv) == 1
    err = capsys.readouterr().err
    assert "Q002: source_excerpt not found in the text of ['50/1/b'] in force on 2025-06-01" in err


@pytest.mark.db
def test_export_with_texts_in_force(
    monkeypatch: pytest.MonkeyPatch,
    seeded_conn: psycopg.Connection,
    dataset_path: Path,
    tmp_path: Path,
) -> None:
    @contextmanager
    def open_db() -> Iterator[psycopg.Connection]:
        yield seeded_conn

    monkeypatch.setattr(eval_cli, "_open_db", open_db)
    path = tmp_path / "review.xlsx"
    assert cli_main(["export-xlsx", str(path), "--dataset", str(dataset_path)]) == 0
    ws = load_workbook(path)[SHEET_REVIEW]
    column = next(cell.column for cell in ws[1] if cell.value == "Testo vigente")
    assert ws.cell(row=2, column=column).value.startswith("[Art. 50, comma 1]\n1. ")
    assert ws.cell(row=4, column=column).value in (None, "")  # unanswerable: no refs


def test_database_unreachable(
    monkeypatch: pytest.MonkeyPatch, dataset_path: Path, capsys: pytest.CaptureFixture
) -> None:
    @contextmanager
    def open_db() -> Iterator[None]:
        raise psycopg.OperationalError("connection refused\ndetails")
        yield

    monkeypatch.setattr(eval_cli, "_open_db", open_db)
    assert cli_main(["validate", "--db", "--dataset", str(dataset_path)]) == 1
    assert "Database not reachable: connection refused" in capsys.readouterr().err


def test_retrieval_command(
    tmp_path: Path, dataset_path: Path, stub_backends: list[str], capsys: pytest.CaptureFixture
) -> None:
    out = tmp_path / "results"
    args = ["retrieval", "--dataset", str(dataset_path), "--all", "--out", str(out)]
    assert cli_main([*args, "--mode", "lexical", "--mode", "hybrid", "--k", "1,5"]) == 0
    assert stub_backends == ["lexical", "hybrid"]
    (path,) = out.glob("*_retrieval.json")
    results = json.loads(path.read_text())
    assert results["config"]["modes"] == ["lexical", "hybrid"]
    assert results["config"]["retrieval_mode"] == "hybrid"
    assert results["dataset"]["selection"] == "all"
    assert not (out / "latest.json").exists()
    printed = capsys.readouterr().out
    assert "retrieval [lexical] n=3 recall@1=0.167 recall@5=0.667 mrr=0.500" in printed


def test_retrieval_command_publish(
    tmp_path: Path, dataset_path: Path, stub_backends: list[str]
) -> None:
    out = tmp_path / "results"
    args = ["retrieval", "--dataset", str(dataset_path), "--all", "--out", str(out)]
    assert cli_main([*args, "--mode", "lexical", "--publish"]) == 0
    latest = json.loads((out / "latest.json").read_text())
    assert latest["retrieval_by_mode"].keys() == {"lexical"}


def test_answers_command(
    tmp_path: Path, dataset_path: Path, stub_backends: list[str], capsys: pytest.CaptureFixture
) -> None:
    out = tmp_path / "results"
    argv = ["answers", "--dataset", str(dataset_path), "--out", str(out), "--ids", "Q001,Q002"]
    assert cli_main(argv) == 0
    results = json.loads(_run_file(out).read_text())
    assert [item["id"] for item in results["items"]] == ["Q001", "Q002"]
    assert results["dataset"]["selection"] == "verified"
    assert results["config"]["provider"] == "fake" and results["config"]["judge"] == "overlap"
    assert results["retrieval"]["overall"]["n"] == 2
    assert _run_file(out).with_suffix(".md").read_text().startswith("# Evaluation report")
    assert not (out / "latest.json").exists() and not (out / "latest.md").exists()
    assert "answers n=2 correctness=0.500" in capsys.readouterr().out


def test_answers_falls_back_to_fake_without_key(
    tmp_path: Path, dataset_path: Path, stub_backends: list[str], capsys: pytest.CaptureFixture
) -> None:
    argv = ["run", "--dataset", str(dataset_path), "--out", str(tmp_path), "--provider", "groq"]
    assert cli_main(argv) == 0
    assert "No API key for groq: answering with the fake provider." in capsys.readouterr().err
    assert not (tmp_path / "latest.json").exists()


def test_answers_publish_refused_with_fake_provider(
    tmp_path: Path, dataset_path: Path, stub_backends: list[str], capsys: pytest.CaptureFixture
) -> None:
    out = tmp_path / "results"
    argv = ["answers", "--dataset", str(dataset_path), "--out", str(out), "--publish"]
    assert cli_main([*argv, "--provider", "anthropic"]) == 2
    err = capsys.readouterr().err
    assert "Refusing to publish: the answers and the judge would use the fake provider" in err
    assert not out.exists()  # refused before running anything


def test_answers_publish_refused_with_fake_judge(
    tmp_path: Path,
    dataset_path: Path,
    stub_backends: list[str],
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture,
) -> None:
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    get_settings.cache_clear()
    argv = ["answers", "--dataset", str(dataset_path), "--out", str(tmp_path), "--publish"]
    assert cli_main([*argv, "--provider", "anthropic"]) == 2
    assert "Refusing to publish: the judge would use the fake provider" in capsys.readouterr().err
    assert not (tmp_path / "latest.json").exists()


def test_answers_publish(
    tmp_path: Path,
    dataset_path: Path,
    stub_backends: list[str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    get_settings.cache_clear()
    monkeypatch.setattr(eval_cli, "get_judge", lambda settings: OverlapJudge())  # no network
    argv = ["answers", "--dataset", str(dataset_path), "--out", str(tmp_path), "--publish"]
    assert cli_main([*argv, "--provider", "anthropic", "--judge-provider", "anthropic"]) == 0
    latest = json.loads((tmp_path / "latest.json").read_text())
    assert latest["config"]["provider"] == "anthropic"
    assert (tmp_path / "latest.md").exists()


def test_answers_retrieve_like_production(
    tmp_path: Path,
    dataset_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    no_db: None,
    stub_retriever: Callable,
    items: list[EvalItem],
) -> None:
    """Answers get search(question, date, k, mode) as in the app; metrics make their own calls."""
    calls: list[tuple[str, str, int]] = []

    def search_retriever(mode: str, conn: object, settings: object) -> Callable:
        def retrieve(question: str, at_date: date, k: int) -> list[Hit]:
            calls.append((mode, question, k))
            return stub_retriever(question, at_date, k)

        return retrieve

    def answer_backend(settings, provider_name, model, retrieve, *, k, mode, delay_s):
        def answer(question: str, at_date: date) -> Answer:
            hits = retrieve(question, at_date, k)
            return Answer(
                question=question,
                at_date=at_date,
                text="-",
                sources=hits,
                abstained=True,
                model="fake",
                provider="fake",
            )

        return answer, provider_name, "fake"

    monkeypatch.setattr(eval_cli, "_search_retriever", search_retriever)
    monkeypatch.setattr(eval_cli, "_answer_backend", answer_backend)
    argv = ["answers", "--dataset", str(dataset_path), "--all", "--out", str(tmp_path)]
    assert cli_main([*argv, "--k", "2", "--mode", "lexical"]) == 0

    by_question = {item.question: item for item in items}
    answer_calls = [(by_question[q].id, k) for mode, q, k in calls if k == 2]
    metric_calls = [(by_question[q].id, k) for mode, q, k in calls if k != 2]
    assert {mode for mode, _, _ in calls} == {"lexical"}
    assert answer_calls == [("Q001", 2), ("Q002", 2), ("Q003", 2), ("Q004", 2)]
    assert metric_calls == [("Q001", 10), ("Q002", 10), ("Q004", 10)]  # Q003 has no refs
    records = {r["id"]: r for r in json.loads(_run_file(tmp_path).read_text())["items"]}
    assert records["Q001"]["retrieved"] == ["50/2", "50/1"]


def test_nothing_selected(dataset_path: Path, capsys: pytest.CaptureFixture) -> None:
    assert cli_main(["answers", "--dataset", str(dataset_path), "--ids", "Q003"]) == 1
    assert "add --all for drafts" in capsys.readouterr().err


def test_report_and_compare(
    tmp_path: Path, dataset_path: Path, stub_backends: list[str], capsys: pytest.CaptureFixture
) -> None:
    assert cli_main(["answers", "--dataset", str(dataset_path), "--out", str(tmp_path)]) == 0
    latest = _run_file(tmp_path)
    capsys.readouterr()
    assert cli_main(["report", str(latest)]) == 0
    assert capsys.readouterr().out.startswith("# Evaluation report")
    output = tmp_path / "report.md"
    assert cli_main(["report", str(latest), "--output", str(output)]) == 0
    assert output.read_text().startswith("# Evaluation report")
    assert cli_main(["compare", str(latest), str(latest)]) == 0
    assert capsys.readouterr().out.count("| fake | fake |") == 2


def test_export_and_import(
    tmp_path: Path,
    dataset_path: Path,
    stub_backends: list[str],
    items: list[EvalItem],
    capsys: pytest.CaptureFixture,
) -> None:
    assert cli_main(["answers", "--dataset", str(dataset_path), "--out", str(tmp_path)]) == 0
    workbook_path = tmp_path / "review.xlsx"
    argv = ["export-xlsx", str(workbook_path), "--dataset", str(dataset_path), "--no-texts"]
    assert cli_main([*argv, "--with-answers", str(_run_file(tmp_path))]) == 0
    workbook = load_workbook(workbook_path)
    ws = workbook[SHEET_REVIEW]
    headers = {cell.value: cell.column for cell in ws[1]}
    assert "Testo vigente" not in headers
    assert ws.cell(row=2, column=headers["Riferimenti citati"]).value == "50/1"
    ws.cell(row=4, column=headers["Esito"], value="OK")
    workbook.save(workbook_path)

    import_argv = ["import-xlsx", str(workbook_path), "--dataset", str(dataset_path)]
    assert cli_main([*import_argv, "--dry-run"]) == 0
    assert "Q003: bozza -> verificata" in capsys.readouterr().out
    assert load_dataset(dataset_path) == items
    assert cli_main(import_argv) == 0
    assert "1 question(s) updated" in capsys.readouterr().out
    assert load_dataset(dataset_path)[2].status == "verificata"


def test_export_without_database(
    tmp_path: Path, dataset_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    @contextmanager
    def open_db() -> Iterator[None]:
        raise psycopg.OperationalError("connection refused")
        yield

    monkeypatch.setattr(eval_cli, "_open_db", open_db)
    path = tmp_path / "review.xlsx"
    assert (
        cli_main(["export-xlsx", str(path), "--dataset", str(dataset_path), "--ids", "Q001"]) == 0
    )
    ws = load_workbook(path)[SHEET_REVIEW]
    assert ws.max_row == 2
    assert "Testo vigente" not in [cell.value for cell in ws[1]]


def test_usage_errors() -> None:
    with pytest.raises(SystemExit) as error:
        cli_main([])
    assert error.value.code == 2
    with pytest.raises(SystemExit):
        cli_main(["retrieval", "--k", "0,5"])


def test_answers_with_the_real_pipeline(
    tmp_path: Path,
    dataset_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    no_db: None,
    stub_retriever: Callable,
) -> None:
    """End to end with answer_question and the FakeProvider; only retrieval is stubbed."""
    try:
        from appalti_qa.answer import answer_question  # noqa: F401
        from appalti_qa.llm import FakeProvider  # noqa: F401
    except ImportError:
        pytest.skip("answer/llm modules not available")
    monkeypatch.setattr(eval_cli, "_search_retriever", lambda mode, conn, settings: stub_retriever)
    argv = ["answers", "--dataset", str(dataset_path), "--all", "--out", str(tmp_path)]
    assert cli_main(argv) == 0
    results = json.loads(_run_file(tmp_path).read_text())
    assert results["config"]["provider"] == "fake"
    records = {item["id"]: item for item in results["items"]}
    assert "error" not in records["Q001"]
    assert records["Q001"]["cited"] == ["50/2", "50/1"]  # the fake quotes the top two sources
    assert records["Q003"]["abstained"] is True  # no sources retrieved
    assert results["answer"]["overall"]["abstention_accuracy"] == 1.0
